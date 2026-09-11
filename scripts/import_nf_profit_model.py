"""Compile a reviewable NF cost model; never modify the workbook or SKU mappings."""

from __future__ import annotations

import argparse
import hashlib
import json
from io import BytesIO
from decimal import Decimal
from pathlib import Path
from typing import Any
from zipfile import ZipFile
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator

from takealot_ops.nf_profit import calculate_cells, compile_mapping_authority


ROOT = Path(__file__).resolve().parents[1]
STORE_CODES = {"101": "store-02", "102": "current", "103": "store-03",
               "c1": "store-04", "c2": "store-05", "c3": "store-06"}


def numeric(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def compile_workbook(source: Path) -> dict[str, Any]:
    expected = json.loads((ROOT / "config/nf_profit_formulas.json").read_text(encoding="utf-8"))
    # WPS adds an unsupported validation id. Drop UI validations in memory only;
    # values, formulas, cached results and the original source remain untouched.
    compatible = BytesIO()
    with ZipFile(source) as original, ZipFile(compatible, "w") as target:
        for info in original.infolist():
            data = original.read(info.filename)
            if info.filename.startswith("xl/worksheets/sheet") and info.filename.endswith(".xml"):
                tree = ElementTree.fromstring(data)
                for child in list(tree):
                    if child.tag.endswith("}dataValidations"):
                        tree.remove(child)
                data = ElementTree.tostring(tree, encoding="utf-8")
            target.writestr(info, data)
    cached = load_workbook(BytesIO(compatible.getvalue()), read_only=True, data_only=True)
    formulas = load_workbook(BytesIO(compatible.getvalue()), read_only=True, data_only=False)
    base = cached["基数"]
    base_rows = list(base.iter_rows(max_row=31, max_col=17, values_only=True))
    categories = {str(r[0]): r[1] for r in base_rows[1:] if r[0]}
    constants = {"vat": base_rows[1][2], "outbound_rmb_kg": base_rows[1][5],
                 "outbound_factor": base_rows[1][7], "delivery_tiers": []}
    # Preserve the IFS order, including its omitted 130000 cm3 / 70 kg branch.
    for row in range(1, 5):
        for offset in (9, 11, 13, 15):
            if row == 2 and offset == 15:
                continue
            constants["delivery_tiers"].append(
                [base_rows[row][8], base_rows[row][offset], base_rows[row][offset + 1]])
    dimensions: dict[str, tuple[int, tuple[Any, ...]]] = {}
    for rownum, rowvalues in enumerate(cached["单件CBM 单重"].iter_rows(
            max_col=24, values_only=True), 1):
        if rowvalues[0]:
            dimensions.setdefault(str(rowvalues[0]).casefold(), (rownum, rowvalues))
    costs: dict[str, tuple[int, tuple[Any, ...]]] = {}
    cost_rows = list(cached["成本&在库统计"].iter_rows(max_col=8, values_only=True))
    authority = compile_mapping_authority(cost_rows)
    for rownum, rowvalues in enumerate(cost_rows, 1):
        if rowvalues[1]:
            costs.setdefault(str(rowvalues[1]).casefold(), (rownum, rowvalues))
    profiles = []
    sheet = cached["利润计算表"]
    raw = formulas["利润计算表"]
    for rownum, (values, formula_cells) in enumerate(zip(
            sheet.iter_rows(max_col=24, values_only=True),
            raw.iter_rows(max_col=24, values_only=True), strict=True), 1):
        if rownum == 1 or not values[1]:
            continue
        v = dict(zip("ABCDEFGHIJKLMNOPQRSTUVWX", values, strict=True))
        issues: list[str] = []
        overrides = {}
        for col in "FGHIJKLMNPQRSTUVWX":
            formula = formula_cells[ord(col) - 65]
            if hasattr(formula, "text"):
                formula = formula.text
            if isinstance(formula, str) and formula.startswith("="):
                translated = (Translator(expected[col], origin=f"{col}2")
                              .translate_formula(f"{col}{rownum}")) if expected[col] else None
                if formula != translated:
                    issues.append(f"利润计算表 {col}{rownum} 使用异常或未支持的引用")
                if not numeric(v[col]):
                    issues.append(f"利润计算表 {col}{rownum} 缺少有效数值或包含公式错误")
            elif col in "TVWX":
                issues.append(f"利润计算表 {col}{rownum} 是手填结果或缺失公式")
            elif not numeric(v[col]):
                issues.append(f"利润计算表 {col}{rownum} 缺少有效数值")
            elif col in "GHJKLMNP":
                overrides[col] = v[col]
        inputs = {col: v[col] for col in "FIQRSUO"}
        inputs["commission_rate"] = categories.get(str(v["E"]), 0 if "K" in overrides else None)
        if inputs["commission_rate"] is None:
            issues.append(f"利润计算表 E{rownum} 缺少可匹配的类目佣金")
        refs = {}
        dim = dimensions.get(str(v["B"]).casefold())
        if dim:
            dimrow, dimvalues = dim
            refs["dimensions"] = f"单件CBM 单重!G{dimrow}:X{dimrow}"
            # S is authoritative: some rows deliberately use a manually rounded O.
            if (not numeric(dimvalues[18]) or dimvalues[18] <= 0
                    or not numeric(v["F"]) or abs(dimvalues[18] - v["F"]) > 1e-9):
                issues.append(f"单件CBM 单重 S{dimrow} 缺少有效分摊体积或与利润行不一致")
            if (not numeric(dimvalues[23]) or dimvalues[23] <= 0
                    or not numeric(v["I"]) or abs(dimvalues[23] - v["I"]) > 1e-9):
                issues.append(f"单件CBM 单重 X{dimrow} 缺少有效单重或与利润行不一致")
        else:
            issues.append(f"利润计算表 B{rownum} 无对应箱规")
        cost = costs.get(str(v["B"]).casefold())
        cost_formula = formula_cells[20]
        if hasattr(cost_formula, "text"):
            cost_formula = cost_formula.text
        uses_cost_lookup = isinstance(cost_formula, str) and cost_formula.startswith("=")
        if cost:
            costrow, costvalues = cost
            refs["cost"] = f"成本&在库统计!F{costrow}:H{costrow}"
            if uses_cost_lookup:
                if (not numeric(v["U"]) or not numeric(costvalues[7])
                        or abs(v["U"] - costvalues[7]) > 1e-7):
                    issues.append(f"利润计算表 U{rownum} 与来源成本不一致")
        elif uses_cost_lookup:
            issues.append(f"利润计算表 U{rownum} 无对应成本")
        stores = str(v["A"] or "").lower().replace("，", ",").split(",")
        profile = {"row": rownum, "company_sku": str(v["B"]).strip(),
                   "stores": sorted({STORE_CODES[s.strip()] for s in stores if s.strip() in STORE_CODES}),
                   "inputs": inputs, "constants": constants, "overrides": overrides,
                   "references": refs, "issues": issues, "reference_price_zar": v["D"]}
        profile["reference_results"] = {col: v[col] for col in "JKLMNOPTVWX"}
        if not issues:
            try:
                computed = calculate_cells(profile, Decimal(str(v["D"])))
                mismatches = [col for col, value in computed.items()
                              if not numeric(v[col]) or abs(float(value) - v[col]) > 1e-7]
                if mismatches:
                    issues.append(f"利润计算表 第{rownum}行缓存结果与公式复算不一致：" + ",".join(mismatches))
            except (ValueError, ArithmeticError, KeyError, TypeError) as exc:
                issues.append(f"利润计算表 第{rownum}行：{exc}")
        profiles.append(profile)
    cached.close()
    formulas.close()
    return {"schema_version": 2,
            "source": {"file": source.name, "sheet": "利润计算表",
                       "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
                       "reference_sheets": ["基数", "单件CBM 单重", "成本&在库统计"]},
            "mapping_authority": authority, "profiles": profiles}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("--output", type=Path, default=ROOT / "config/nf_profit_model.json")
    args = parser.parse_args()
    model = compile_workbook(args.source)
    args.output.write_text(json.dumps(model, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"profiles": len(model["profiles"]),
                      "verified": sum(not p["issues"] for p in model["profiles"]),
                      "output": str(args.output)}, ensure_ascii=False))
